"""The Excel side of migration requests: a template to fill in, and a reader.

Only the fields a developer owns appear here. Approvals, QA and production
results are deliberately absent — those are role-gated events with an audit
trail, and a spreadsheet upload is not a way to assert that an approver approved
something.

The template carries real dropdowns (Excel data validation) driven by a Lists
sheet, so the permitted values travel with the file instead of living only in
the web form. Repo and track lead depend on the microservice, which Excel cannot
express as a cascade, so the Lists sheet spells out the mapping for lookup and
the server validates the pairing on upload.
"""

from __future__ import annotations

import io
from typing import Any

from .config import MigrationConfig
from .migrations import (
    BY_KEY, REQUEST_KEYS, YESNO, MigrationError, clean_requestor, validate, _clean,
)
from .spreadsheet import SheetError, parse_columns

# Longer, more specific aliases first — the reader also falls back to a
# "contains" match, so a header like "Commit Hash (if code change)" still lands.
SHEET_COLUMNS: dict[str, list[str]] = {
    "release": ["rel#", "rel no", "release no", "release number", "release", "rel"],
    "migration_path": ["migration path", "movement path", "path"],
    "microservice": ["micro service name", "microservice name", "micro service",
                     "microservice", "service name", "service"],
    "repo_name": ["repo name", "repository name", "repository", "repo"],
    "track_lead": ["track lead name", "track lead", "lead name", "lead"],
    "change_requestor": ["change requestor", "change requester", "requested by",
                         "requestor", "requester"],
    "reason": ["reason for movement", "reason for change", "reason"],
    "change_description": ["change description", "change desc", "description"],
    "code_image_change": ["code & image change?", "code & image change",
                          "code and image change", "code image change", "code change"],
    "commit_hash": ["commit hash", "commit id", "commit sha", "commit", "sha"],
    "env_change": ["environment change", "env change"],
    "env_secret": ["env secret details", "environment secret", "env secret", "secret details"],
    "ddl_dml": ["ddl/dml", "ddl dml", "ddl", "dml"],
    "db_script_path": ["db script path", "database script path", "db script", "script path"],
}

MAX_UPLOAD_ROWS = 200


def parse(filename: str, content: bytes) -> dict[str, Any]:
    return parse_columns(
        filename,
        content,
        columns=SHEET_COLUMNS,
        required="microservice",
        hint=(
            "The sheet needs at least a “Micro Service Name” column. "
            "Download the template from this tab to see the expected shape."
        ),
    )


def check_rows(rows: list[dict[str, Any]], cfg: MigrationConfig) -> list[dict[str, Any]]:
    """Validate every row. Returns one result per row, in sheet order.

    Nothing is written here — the caller decides whether to import, so the UI can
    show what would happen before it happens.
    """
    if len(rows) > MAX_UPLOAD_ROWS:
        raise MigrationError(
            f"{len(rows)} rows is more than the {MAX_UPLOAD_ROWS} this will import at once."
        )

    out: list[dict[str, Any]] = []
    for raw in rows:
        values: dict[str, str] = {}
        for key in REQUEST_KEYS:
            if key in raw:
                values[key] = _clean(BY_KEY[key], raw[key])
        if values.get("change_requestor"):
            # A sheet usually carries just the employee number.
            values["change_requestor"] = clean_requestor(values["change_requestor"], cfg)
        for key in REQUEST_KEYS:
            spec = BY_KEY[key]
            if spec.default and not values.get(key):
                values[key] = spec.default

        problems = validate(values, cfg)
        out.append({
            "row": raw.get("_row"),
            "values": values,
            "ok": not problems,
            "problems": {BY_KEY[k].label: v for k, v in problems.items()},
        })
    return out


def build_template(cfg: MigrationConfig) -> bytes:
    """A workbook with the request columns, real dropdowns, and a lookup sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Migration Requests"
    lists = book.create_sheet("Lists")

    keys = list(REQUEST_KEYS)
    headers = [BY_KEY[k].label for k in keys]
    sheet.append(headers)

    header_fill = PatternFill("solid", fgColor="0A7568")
    amber_fill = PatternFill("solid", fgColor="B45309")
    for index, key in enumerate(keys, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        # The risk flags are the ones an approver will look at first.
        cell.fill = amber_fill if BY_KEY[key].amber_when else header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    # --- the Lists sheet backing the dropdowns ---------------------------- #
    named: dict[str, str] = {}  # field key -> "Lists!$A$2:$A$9"
    column = 1

    def put(title: str, values: list[str]) -> str:
        nonlocal column
        letter = get_column_letter(column)
        lists.cell(row=1, column=column, value=title).font = Font(bold=True)
        for offset, value in enumerate(values, start=2):
            lists.cell(row=offset, column=column, value=value)
        lists.column_dimensions[letter].width = max(16, min(40, len(title) + 6))
        column += 1
        # An empty list would make an unusable validation, so hand back "".
        return f"Lists!${letter}$2:${letter}${len(values) + 1}" if values else ""

    named["release"] = put("Rel#", cfg.releases)
    named["migration_path"] = put("Migration Path", cfg.migration_paths)
    named["microservice"] = put("Micro Service", cfg.service_names)
    # Employee labels where staff are configured, the old plain list otherwise.
    named["change_requestor"] = put(
        "Change Requestor", cfg.employee_labels or cfg.change_requestors
    )
    yesno_ref = put("Yes / No", list(YESNO))

    # Repo and track lead cascade off the microservice, which Excel validation
    # cannot express — so the mapping is written out to be read, not enforced.
    lists.cell(row=1, column=column, value="Micro Service").font = Font(bold=True)
    lists.cell(row=1, column=column + 1, value="Repo Name").font = Font(bold=True)
    lists.cell(row=1, column=column + 2, value="Track Lead").font = Font(bold=True)
    for letter, width in zip(
        (get_column_letter(column), get_column_letter(column + 1), get_column_letter(column + 2)),
        (24, 34, 24),
    ):
        lists.column_dimensions[letter].width = width
    line = 2
    for service in cfg.microservices:
        pairs = max(len(service.repos), len(service.track_leads), 1)
        for index in range(pairs):
            lists.cell(row=line, column=column, value=service.name)
            if index < len(service.repos):
                lists.cell(row=line, column=column + 1, value=service.repos[index])
            if index < len(service.track_leads):
                lists.cell(row=line, column=column + 2, value=service.track_leads[index])
            line += 1

    # --- attach the validations ------------------------------------------- #
    last = 500
    for index, key in enumerate(keys, start=1):
        spec = BY_KEY[key]
        ref = yesno_ref if spec.kind == "yesno" else named.get(key, "")
        if not ref:
            continue
        validation = DataValidation(type="list", formula1=f"={ref}", allow_blank=True)
        validation.error = f"Pick a value from the {spec.label} list."
        validation.errorTitle = "Not a permitted value"
        sheet.add_data_validation(validation)
        letter = get_column_letter(index)
        validation.add(f"{letter}2:{letter}{last}")

    widths = {"reason": 40, "change_description": 46, "commit_hash": 24, "db_script_path": 34}
    for index, key in enumerate(keys, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(
            key, max(14, min(30, len(BY_KEY[key].label) + 4))
        )
    sheet.freeze_panes = "A2"

    # One worked example, using real configured values where they exist.
    service = cfg.microservices[0] if cfg.microservices else None
    example = {
        "release": cfg.releases[0] if cfg.releases else "",
        "migration_path": cfg.migration_paths[0] if cfg.migration_paths else "",
        "microservice": service.name if service else "",
        "repo_name": service.repos[0] if service and service.repos else "",
        "track_lead": service.track_leads[0] if service and service.track_leads else "",
        "change_requestor": (cfg.employee_labels or cfg.change_requestors or [""])[0],
        "reason": "Defect fix agreed in the 09:30 triage call",
        "change_description": "Corrects the rounding on the settlement total",
        "code_image_change": "Yes",
        "commit_hash": "0000000",
        "env_change": "No",
        "env_secret": "No",
        "ddl_dml": "No",
        "db_script_path": "",
    }
    sheet.append([example.get(k, "") for k in keys])
    for cell in sheet[2]:
        cell.font = Font(italic=True, color="6B7280")

    note = sheet.cell(row=4, column=1)
    note.value = (
        "Row 2 is an example — replace or delete it. "
        "Commit Hash is required when Code & Image Change is Yes; "
        "DB Script Path is required when DDL/DML is Yes. "
        "See the Lists sheet for which repo and track lead belong to each microservice."
    )
    note.font = Font(italic=True, color="6B7280")

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


__all__ = ["SHEET_COLUMNS", "MAX_UPLOAD_ROWS", "parse", "check_rows", "build_template", "SheetError"]
