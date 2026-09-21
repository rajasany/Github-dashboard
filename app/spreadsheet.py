"""Read an uploaded .xlsx / .csv into rows, tolerating real-world layouts.

Spreadsheets people actually keep are rarely a clean header-in-A1 grid: there is
often a title line, a blank row, merged notes. So the header row is *found*
rather than assumed, and columns are matched by alias rather than position — a
sheet with "Repository", "Service" and "Branch" works as well as one with
"repo", "folder", "branch".
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

MAX_ROWS = 500
# How far down to look for the header before giving up.
HEADER_SCAN_ROWS = 15

# Longer aliases first: "commit hash" must win over "commit".
COLUMNS: dict[str, list[str]] = {
    "repo": ["repository name", "repo name", "repository", "repo", "project name", "project", "git repo"],
    "folder": ["service / folder", "service/folder", "folder path", "sub folder", "subfolder",
               "directory", "folder", "service", "module", "component", "path", "area"],
    "branch": ["branch name", "branch", "ref", "target branch"],
    "commit": ["commit hash", "commit id", "commit sha", "commit", "hash", "sha", "revision"],
    "tag": ["proposed tag", "tag name", "tag"],
    "note": ["notes", "note", "comment", "comments", "remarks", "description"],
}


class SheetError(Exception):
    pass


def _norm(value: Any) -> str:
    return re.sub(r"[\s_\-]+", " ", str(value or "").strip().lower())


def _match_column(header: str, columns: dict[str, list[str]]) -> str | None:
    text = _norm(header)
    if not text:
        return None
    for field, aliases in columns.items():
        for alias in aliases:
            if text == alias:
                return field
    # Fall back to a contains match, longest alias first, so "Repo (owner/name)"
    # still resolves without matching something shorter by accident.
    best: tuple[int, str] | None = None
    for field, aliases in columns.items():
        for alias in aliases:
            if alias in text and (best is None or len(alias) > best[0]):
                best = (len(alias), field)
    return best[1] if best else None


def _grid_from_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise SheetError(f"That file could not be read as a spreadsheet: {exc}") from exc
    sheet = book.worksheets[0]
    grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    book.close()
    return grid


def _grid_from_csv(content: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SheetError("That file is not readable as text.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def parse_sheet(filename: str, content: bytes) -> dict[str, Any]:
    """Read a release-plan sheet. Return {columns, rows, header_row, sheet_rows, ignored}."""
    return parse_columns(
        filename,
        content,
        columns=COLUMNS,
        required="repo",
        hint=(
            "The sheet needs at least a column named “Repository” (and usually "
            "“Folder” and “Branch”). Download the template from this tab to see "
            "the expected shape."
        ),
    )


def parse_columns(
    filename: str,
    content: bytes,
    *,
    columns: dict[str, list[str]],
    required: str,
    hint: str,
) -> dict[str, Any]:
    """The generic reader: find the header row, map columns by alias, return rows.

    `required` names the column a row must have a value in to count as a row at
    all, and which the header must contain for the sheet to be usable.
    """
    if not content:
        raise SheetError("That file is empty.")

    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        grid = _grid_from_xlsx(content)
    elif name.endswith((".csv", ".tsv", ".txt")):
        grid = _grid_from_csv(content)
    elif name.endswith(".xls"):
        raise SheetError(
            "The old .xls format is not supported. Re-save it as .xlsx or .csv."
        )
    else:
        raise SheetError("Upload a .xlsx or .csv file.")

    if not grid:
        raise SheetError("That spreadsheet has no rows.")

    # Find the header: the row that resolves the most known columns.
    best_index, best_map, best_score = -1, {}, 0
    for index, row in enumerate(grid[:HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        for position, cell in enumerate(row):
            field = _match_column(cell, columns)
            # First column to claim a field keeps it.
            if field and field not in mapping.values():
                mapping[position] = field
        score = len(mapping)
        if score > best_score:
            best_index, best_map, best_score = index, mapping, score

    if best_score == 0 or required not in best_map.values():
        raise SheetError(f"No usable header row found. {hint}")

    rows: list[dict[str, Any]] = []
    ignored = 0
    for raw in grid[best_index + 1 :]:
        record = {"_row": best_index + 2 + len(rows) + ignored}
        for position, field in best_map.items():
            if position < len(raw):
                value = raw[position]
                record[field] = str(value).strip() if value is not None else ""
        if not (record.get(required) or "").strip():
            ignored += 1
            continue
        rows.append(record)
        if len(rows) >= MAX_ROWS:
            break

    if not rows:
        label = required.replace("_", " ")
        raise SheetError(f"The header was found, but no row had a {label} in it.")

    return {
        "columns": sorted(set(best_map.values())),
        "rows": rows,
        "header_row": best_index + 1,
        "sheet_rows": len(grid),
        "ignored": ignored,
    }


def build_template() -> bytes:
    """A starter workbook, so the expected shape is discoverable, not guessed."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Releases"

    headers = ["Repository", "Folder", "Branch", "Notes"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0A7568")

    sheet.append(["rajasany/Insurance", "src", "main", "Folder and Branch are optional"])
    sheet.append(["rajasany/Github-dashboard", "app", "", "Blank branch uses the default"])
    sheet.append(["rajasany/Guardrails", "", "main", "Blank folder covers the whole repo"])

    for column, width in zip("ABCD", (34, 22, 16, 44)):
        sheet.column_dimensions[column].width = width

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
